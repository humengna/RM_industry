# coding: utf-8
"""
把 Excel 版的指数历史成分股转成回测用的缓存 CSV。

输入：一张带「生效日期 / 失效日期 / 证券代码」的表（列名可有别名，见 COLUMN_ALIASES），
      每行代表"某证券在某个有效区间内属于该指数"。
输出：两列 CSV（date,stock），每个生效日期一个快照，快照里列出当时的全部成分股。
      这正是 hs300_ma_divergence/constituents.py 读取的格式。

用法：
    python tools/import_constituents_xlsx.py 沪深300历史成分.xlsx
    python tools/import_constituents_xlsx.py in.xlsx -o data/hs300_constituents.csv --sheet 历史成分
"""

import argparse
import os
import re
import sys

COLUMN_ALIASES = {
    'effective': ('生效日期', '纳入日期', '调入日期', '开始日期', 'start_date', 'in_date'),
    'expire': ('失效日期', '剔除日期', '调出日期', '结束日期', 'end_date', 'out_date'),
    'code': ('证券代码', '成分券代码', '股票代码', '代码', 'code', 'stock', 'symbol'),
    'name': ('证券简称', '成分券名称', '股票名称', '名称', 'name'),
}

# 无后缀代码按前缀补市场
SUFFIX_BY_PREFIX = (
    (('60', '68', '51', '58', '56', '50', '11', '90'), 'SH'),
    (('00', '30', '20', '15', '16', '18', '12'), 'SZ'),
    (('43', '83', '87', '88', '92'), 'BJ'),
)

CODE_PATTERN = re.compile(r'^(\d{6})\.(SH|SZ|BJ)$')


def import_openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise SystemExit('需要 openpyxl 才能读 xlsx：pip install openpyxl')
    return openpyxl


def normalize_code(value):
    """
    统一成 QMT 的 '600000.SH' 形式。

    支持 '600000.SH' / '600000' / 600000 / 'SH600000' / 'sh.600000'。
    认不出来的返回 None。
    """
    if value is None:
        return None

    text = str(value).strip().upper().replace(' ', '')
    if not text:
        return None

    match = CODE_PATTERN.match(text)
    if match:
        return '%s.%s' % match.groups()

    # SH600000 / SH.600000
    match = re.match(r'^(SH|SZ|BJ)\.?(\d{6})$', text)
    if match:
        return '%s.%s' % (match.group(2), match.group(1))

    # 纯数字（Excel 常把 000001 存成数字 1）
    digits = re.sub(r'\D', '', text)
    if not digits:
        return None
    digits = digits.zfill(6)
    if len(digits) != 6:
        return None

    for prefixes, suffix in SUFFIX_BY_PREFIX:
        if digits.startswith(prefixes):
            return '%s.%s' % (digits, suffix)
    return None


def to_date(value):
    """日期单元格 -> 'YYYYMMDD'；空值返回 ''（表示仍然有效）。"""
    if value is None:
        return ''
    if hasattr(value, 'strftime'):
        return value.strftime('%Y%m%d')

    digits = re.sub(r'\D', '', str(value))
    return digits[:8] if len(digits) >= 8 else ''


def find_header(rows):
    """找到表头行，返回 (行号, {字段: 列号})。"""
    for index, row in enumerate(rows):
        mapping = {}
        for column, cell in enumerate(row):
            if cell is None:
                continue
            text = str(cell).strip()
            for field, aliases in COLUMN_ALIASES.items():
                if field not in mapping and text in aliases:
                    mapping[field] = column
        if 'effective' in mapping and 'code' in mapping:
            return index, mapping
    return -1, {}


def pick_sheet(workbook, name=None):
    """优先用指定的 sheet，其次找名字里带"成分"的，最后用第一个。"""
    if name:
        if name not in workbook.sheetnames:
            raise SystemExit('工作簿里没有 sheet：%s（现有：%s）'
                             % (name, ', '.join(workbook.sheetnames)))
        return workbook[name]

    for sheet_name in workbook.sheetnames:
        if '成分' in sheet_name and '明细' not in sheet_name:
            return workbook[sheet_name]
    return workbook[workbook.sheetnames[0]]


def read_intervals(path, sheet_name=None):
    """读出 [(生效日, 失效日, 代码)]，代码已标准化。"""
    openpyxl = import_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = pick_sheet(workbook, sheet_name)

    rows = list(sheet.iter_rows(values_only=True))
    header_index, mapping = find_header(rows)
    if header_index < 0:
        raise SystemExit('没找到表头，需要至少包含「生效日期」和「证券代码」两列')

    intervals = []
    skipped = []
    for row in rows[header_index + 1:]:
        if not row:
            continue
        code = normalize_code(row[mapping['code']] if mapping['code'] < len(row) else None)
        effective = to_date(row[mapping['effective']]
                            if mapping['effective'] < len(row) else None)
        if not code or not effective:
            raw = row[mapping['code']] if mapping['code'] < len(row) else None
            if raw:
                skipped.append(raw)
            continue

        expire = ''
        if 'expire' in mapping and mapping['expire'] < len(row):
            expire = to_date(row[mapping['expire']])
        intervals.append((effective, expire, code))

    return intervals, sheet.title, skipped


def build_snapshots(intervals):
    """
    区间记录 -> {生效日: [成分股]}。

    每个生效日的成分 = 所有"生效日 <= 该日 且（失效日为空 或 失效日 >= 该日）"的证券，
    这样即使表里一只票只写了一条长区间，也能正确出现在每个快照里。
    """
    snapshot_days = sorted({effective for effective, _expire, _code in intervals})
    snapshots = {}
    for day in snapshot_days:
        members = set()
        for effective, expire, code in intervals:
            if effective <= day and (not expire or expire >= day):
                members.add(code)
        snapshots[day] = sorted(members)
    return snapshots


def write_csv(path, snapshots):
    directory = os.path.dirname(os.path.abspath(path))
    if directory and not os.path.isdir(directory):
        os.makedirs(directory)

    with open(path, 'w', encoding='utf-8', newline='') as fp:
        fp.write('date,stock\n')
        for day in sorted(snapshots):
            for stock in snapshots[day]:
                fp.write('%s,%s\n' % (day, stock))
    return path


def describe(snapshots):
    """打印快照概况与每期调入调出，方便肉眼核对。"""
    days = sorted(snapshots)
    print('快照 %d 个：%s ~ %s' % (len(days), days[0], days[-1]))
    print('不同证券 %d 只' % len({s for members in snapshots.values() for s in members}))

    previous = None
    for day in days:
        current = set(snapshots[day])
        if previous is None:
            print('  %s  成分 %d（期初）' % (day, len(current)))
        else:
            print('  %s  成分 %d  调入 %d  调出 %d'
                  % (day, len(current), len(current - previous), len(previous - current)))
        previous = current


def main(argv=None):
    parser = argparse.ArgumentParser(description='Excel 历史成分股 -> 回测缓存 CSV')
    parser.add_argument('xlsx', help='输入的 xlsx 文件')
    parser.add_argument('-o', '--output', default='data/hs300_constituents.csv',
                        help='输出 CSV 路径')
    parser.add_argument('--sheet', default=None, help='指定 sheet 名')
    parser.add_argument('--expect', type=int, default=None,
                        help='每期应有的成分数（如 300），不符会告警')
    args = parser.parse_args(argv)

    intervals, sheet_title, skipped = read_intervals(args.xlsx, args.sheet)
    print('读取 sheet「%s」，有效区间记录 %d 条' % (sheet_title, len(intervals)))
    if skipped:
        print('跳过无法识别的代码 %d 条，例如：%s' % (len(skipped), skipped[:5]))
    if not intervals:
        raise SystemExit('没有读到任何成分股记录')

    snapshots = build_snapshots(intervals)
    describe(snapshots)

    if args.expect:
        odd = [(day, len(members)) for day, members in sorted(snapshots.items())
               if len(members) != args.expect]
        if odd:
            print('注意：以下快照成分数不等于 %d：%s' % (args.expect, odd))

    write_csv(args.output, snapshots)
    print('已写入 %s（%d 行）'
          % (args.output, sum(len(m) for m in snapshots.values()) + 1))
    return 0


if __name__ == '__main__':
    sys.exit(main())
